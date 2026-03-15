import os
import json
from typing import Dict, Any, List

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.shapes.base import BaseShape
except ImportError:
    raise ImportError("The 'python-pptx' library is required. Install it using 'pip install python-pptx'")

def _extract_text_formatting(run) -> Dict[str, Any]:
    """Helper to extract formatting from a text run."""
    font = run.font
    formatting = {
        "text": run.text,
        "is_bold": font.bold if font.bold is not None else False,
        "is_italic": font.italic if font.italic is not None else False,
        "is_underline": font.underline if font.underline is not None else False,
        "font_name": font.name if font.name else "Default",
        "font_size_pt": font.size.pt if font.size else None,
    }
    
    # Try to extract color safely
    try:
        if font.color and font.color.type is not None:
            if hasattr(font.color, 'rgb') and font.color.rgb:
                formatting["color_rgb"] = str(font.color.rgb)
            elif hasattr(font.color, 'theme_color') and font.color.theme_color:
                formatting["theme_color"] = str(font.color.theme_color)
    except Exception:
        pass
        
    return formatting

def _extract_shape_info(shape: BaseShape) -> Dict[str, Any]:
    """Recursively extract information from a single shape or group of shapes."""
    shape_data = {
        "shape_id": shape.shape_id,
        "name": shape.name,
        "type": str(shape.shape_type),
    }

    # Extract spatial information
    try:
        shape_data["position"] = {
            "left_pt": shape.left.pt if hasattr(shape, 'left') and shape.left else None,
            "top_pt": shape.top.pt if hasattr(shape, 'top') and shape.top else None,
            "width_pt": shape.width.pt if hasattr(shape, 'width') and shape.width else None,
            "height_pt": shape.height.pt if hasattr(shape, 'height') and shape.height else None,
            "rotation_deg": shape.rotation if hasattr(shape, 'rotation') and shape.rotation else 0.0,
        }
    except Exception:
        pass

    # Handle Text
    if shape.has_text_frame:
        text_frame = shape.text_frame
        shape_data["text_content"] = text_frame.text
        paragraphs_data = []
        for paragraph in text_frame.paragraphs:
            para_info = {
                "text": paragraph.text,
                "alignment": str(paragraph.alignment) if paragraph.alignment else None,
                "runs": [_extract_text_formatting(run) for run in paragraph.runs if run.text.strip()]
            }
            if para_info["text"].strip():
                paragraphs_data.append(para_info)
        
        shape_data["paragraphs"] = paragraphs_data

    # Handle Pictures
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or shape.shape_type == MSO_SHAPE_TYPE.LINKED_PICTURE:
        shape_data["is_picture"] = True
        try:
            image_blob = shape.image.blob
            shape_data["image_size_bytes"] = len(image_blob)
            shape_data["image_ext"] = shape.image.ext
            shape_data["image_content_type"] = shape.image.content_type
            shape_data["image_filename"] = shape.image.filename
        except Exception:
            shape_data["image_status"] = "Extractable image blob not found or accessible"

    # Handle Group Shapes (Recursion)
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        shape_data["is_group"] = True
        child_shapes = []
        for child_shape in shape.shapes:
            child_shapes.append(_extract_shape_info(child_shape))
        shape_data["group_members"] = child_shapes

    # Handle Tables
    if shape.has_table:
        table = shape.table
        shape_data["is_table"] = True
        shape_data["table_structure"] = {
            "rows": len(table.rows),
            "columns": len(table.columns)
        }
        
        cells_data = []
        for r_idx, row in enumerate(table.rows):
            row_data = []
            for c_idx, cell in enumerate(row.cells):
                cell_text = cell.text_frame.text if cell.text_frame else ""
                row_data.append(cell_text)
            cells_data.append(row_data)
        
        shape_data["table_content"] = cells_data

    return {"shape_details": shape_data}

def extract_slide_context(pptx_path: str, slide_index: int = 0) -> Dict[str, Any]:
    """
    Extracts comprehensive structural and content data from a specified slide in a PowerPoint file.
    
    Args:
        pptx_path: Path to the .pptx file.
        slide_index: The index of the slide to extract (0-indexed).
        
    Returns:
        A dictionary containing the slide dimensions, properties, and all shapes/text/images.
    """
    if not os.path.exists(pptx_path):
        raise FileNotFoundError(f"The file {pptx_path} does not exist.")

    prs = Presentation(pptx_path)
    
    if slide_index < 0 or slide_index >= len(prs.slides):
        raise ValueError(f"Slide index {slide_index} out of range. The presentation has {len(prs.slides)} slides.")

    slide = prs.slides[slide_index]
    
    context = {
        "presentation_file": os.path.basename(pptx_path),
        "slide_index": slide_index,
        "slide_layout_name": slide.slide_layout.name,
        "dimensions_pt": {
            "width": prs.slide_width.pt,
            "height": prs.slide_height.pt
        },
        "shapes": []
    }

    # Iterate over all shapes in the slide and extract details
    for shape in slide.shapes:
        shape_info = _extract_shape_info(shape)
        context["shapes"].append(shape_info)

    return context

def export_slide_context_for_llm(pptx_path: str, slide_index: int = 0) -> str:
    """
    Wrapper function that extracts slide context and returns it as a formatted JSON string
    optimized for injecting into an LLM prompt.
    """
    context_dict = extract_slide_context(pptx_path, slide_index)
    
    # We use indent=2 to ensure the nested structure is easily readable by the LLM
    return json.dumps(context_dict, indent=2, ensure_ascii=False)

if __name__ == "__main__":
    # Example usage for testing purposes:
    # print(export_slide_context_for_llm("example.pptx", 0))
    pass
