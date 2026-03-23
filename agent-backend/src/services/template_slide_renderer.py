"""
Clone slides from user .pptx templates, inject copy, apply design tokens, and
reduce obvious text overlaps (heuristic — not a full layout engine).
"""
from __future__ import annotations

import os
from copy import deepcopy
from typing import Any, Dict, List, Optional, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt
from pptx.oxml.ns import qn


def _hex_to_rgb_color(hex_color: str) -> RGBColor:
    value = (hex_color or "").strip().lstrip("#")
    if len(value) != 6:
        value = "0F172A"
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _clear_slide_shapes(slide) -> None:
    for shape in list(slide.shapes):
        el = shape.element
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def _copy_shapes_from_slide(source_slide, dest_slide) -> None:
    sp_tree = dest_slide.shapes._spTree
    ext_lst = sp_tree.find(qn("p:extLst"))
    for shape in source_slide.shapes:
        new_el = deepcopy(shape.element)
        if ext_lst is not None:
            ext_lst.addprevious(new_el)
        else:
            sp_tree.append(new_el)


def _text_shapes_in_reading_order(slide) -> List[Any]:
    shapes = [s for s in slide.shapes if s.has_text_frame]
    return sorted(shapes, key=lambda s: (s.top, s.left))


def _strip_textframe_to_single_empty(tf) -> None:
    """Remove extra paragraphs; keep one empty paragraph (python-pptx safe)."""
    while len(tf.paragraphs) > 1:
        el = tf.paragraphs[-1]._element
        el.getparent().remove(el)
    tf.paragraphs[0].text = ""


def _set_paragraphs(tf, lines: List[str], font_name: str, size_pt: float, color: RGBColor, bold: bool = False) -> None:
    _strip_textframe_to_single_empty(tf)
    if not lines:
        return
    tf.paragraphs[0].text = lines[0]
    tf.paragraphs[0].level = 0
    if tf.paragraphs[0].runs:
        r = tf.paragraphs[0].runs[0]
        r.font.name = font_name
        r.font.size = Pt(size_pt)
        r.font.bold = bold
        r.font.color.rgb = color
    for line in lines[1:]:
        p = tf.add_paragraph()
        p.text = line
        p.level = 0
        if p.runs:
            r = p.runs[0]
            r.font.name = font_name
            r.font.size = Pt(size_pt)
            r.font.bold = bold
            r.font.color.rgb = color


def _set_simple_text(tf, text: str, font_name: str, size_pt: float, color: RGBColor, bold: bool = False) -> None:
    _set_paragraphs(tf, [text] if text else [""], font_name, size_pt, color, bold)


def _fill_ordered_text_shapes(
    slide,
    title: str,
    subtitle: str,
    bullets: List[str],
    heading_font: str,
    body_font: str,
    accent: RGBColor,
    text_color: RGBColor,
) -> None:
    shapes = _text_shapes_in_reading_order(slide)
    if not shapes:
        return

    if len(shapes) == 1:
        block = title
        if subtitle:
            block = f"{title}\n{subtitle}"
        if bullets:
            block = block + "\n" + "\n".join(f"• {b}" for b in bullets)
        _set_simple_text(shapes[0].text_frame, block.strip(), heading_font, 24, text_color, bold=False)
        if title and shapes[0].text_frame.paragraphs:
            shapes[0].text_frame.paragraphs[0].font.bold = True
            shapes[0].text_frame.paragraphs[0].font.color.rgb = accent
        return

    _set_simple_text(shapes[0].text_frame, title, heading_font, 32, accent, bold=True)
    if len(shapes) >= 2:
        _set_simple_text(shapes[1].text_frame, subtitle, body_font, 18, text_color, bold=False)

    if len(shapes) >= 3:
        body_shapes = shapes[2:]
        if len(body_shapes) == 1:
            _set_paragraphs(body_shapes[0].text_frame, bullets, body_font, 16, text_color, bold=False)
        else:
            n = len(body_shapes)
            buckets: List[List[str]] = [[] for _ in range(n)]
            for i, b in enumerate(bullets):
                buckets[i % n].append(b)
            for shp, bucket in zip(body_shapes, buckets):
                lines = [f"• {x}" for x in bucket] if bucket else [""]
                _set_paragraphs(shp.text_frame, lines, body_font, 15, text_color, bold=False)


def _apply_fonts_colors(slide, heading_font: str, body_font: str, text_color: RGBColor, accent: RGBColor) -> None:
    shapes = _text_shapes_in_reading_order(slide)
    for idx, shp in enumerate(shapes):
        for p in shp.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = heading_font if idx == 0 else body_font
                r.font.color.rgb = accent if idx == 0 else text_color


def _shape_text_rect(shape) -> Tuple[int, int, int, int]:
    return (
        int(shape.left),
        int(shape.top),
        int(shape.left + shape.width),
        int(shape.top + shape.height),
    )


def _rects_overlap(a: Tuple[int, int, int, int], b: Tuple[int, int, int, int]) -> bool:
    return not (a[2] <= b[0] or a[0] >= b[2] or a[3] <= b[1] or a[1] >= b[3])


def _reduce_overlap_fonts(slide, min_pt: float = 9.0) -> None:
    shapes = _text_shapes_in_reading_order(slide)
    if len(shapes) < 2:
        return
    for _ in range(14):
        rects = [_shape_text_rect(s) for s in shapes]
        overlaps = False
        for i in range(len(rects)):
            for j in range(i + 1, len(rects)):
                if _rects_overlap(rects[i], rects[j]):
                    overlaps = True
                    break
            if overlaps:
                break
        if not overlaps:
            break
        for shp in shapes:
            for p in shp.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size:
                        cur = r.font.size.pt
                        if cur > min_pt:
                            r.font.size = Pt(cur - 0.75)


def add_slide_from_template_file(
    target_prs: Presentation,
    template_path: str,
    slide_index: int,
    slots: Dict[str, Any],
    design: Dict[str, Any],
) -> None:
    if not os.path.isfile(template_path):
        raise FileNotFoundError(template_path)

    src_prs = Presentation(template_path)
    if slide_index < 0 or slide_index >= len(src_prs.slides):
        slide_index = 0
    src_slide = src_prs.slides[slide_index]

    layout = target_prs.slide_layouts[6] if len(target_prs.slide_layouts) > 6 else target_prs.slide_layouts[-1]
    dest_slide = target_prs.slides.add_slide(layout)
    _clear_slide_shapes(dest_slide)
    _copy_shapes_from_slide(src_slide, dest_slide)

    css = design.get("css_variables", {})
    token = design.get("design_token", {})
    palette = token.get("palette", {})
    fonts = token.get("fonts", {})

    text_hex = palette.get("text") or css.get("--color-text", "#0F172A")
    accent_hex = palette.get("accent") or css.get("--color-accent", "#2563EB")
    heading_font = fonts.get("heading") or css.get("--font-heading", "Inter")
    body_font = fonts.get("body") or css.get("--font-body", "Inter")

    text_rgb = _hex_to_rgb_color(str(text_hex).replace('"', ""))
    accent_rgb = _hex_to_rgb_color(str(accent_hex).replace('"', ""))

    title = str(slots.get("title", ""))
    subtitle = str(slots.get("subtitle", ""))
    body = slots.get("body", [])
    if isinstance(body, str):
        bullets = [body] if body else []
    else:
        bullets = [str(x) for x in body if str(x).strip()]

    _fill_ordered_text_shapes(
        dest_slide, title, subtitle, bullets, heading_font, body_font, accent_rgb, text_rgb
    )
    _apply_fonts_colors(dest_slide, heading_font, body_font, text_rgb, accent_rgb)
    _reduce_overlap_fonts(dest_slide)


def build_deck_from_templates(rendered: Dict[str, Any], output_path: str) -> str:
    """Build pptx using per-slide template files when paths exist."""
    slides_spec = rendered.get("slides", [])
    if not slides_spec:
        raise ValueError("No slides in rendered_presentation")

    out = Presentation()
    for s in slides_spec:
        tp = s.get("template_pptx_path")
        if tp and os.path.isfile(tp):
            src = Presentation(tp)
            out.slide_width = src.slide_width
            out.slide_height = src.slide_height
            break

    for s in slides_spec:
        tp = s.get("template_pptx_path")
        tidx = int(s.get("template_slide_index", 0))
        slots = s.get("slots", {})
        if tp and os.path.isfile(tp):
            add_slide_from_template_file(out, tp, tidx, slots, rendered)
        else:
            raise ValueError(f"Missing template file for slide {s.get('slide_number')}: {tp}")

    out.save(output_path)
    return output_path
