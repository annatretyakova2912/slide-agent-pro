from __future__ import annotations

import os
from typing import Dict, Any, List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from .template_slide_renderer import add_slide_from_template_file


def _hex_to_rgb(hex_color: str) -> RGBColor:
    value = (hex_color or "").strip().lstrip("#")
    if len(value) != 6:
        value = "0F172A"
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _as_lines(slots: Dict[str, Any]) -> List[str]:
    body = slots.get("body", [])
    if isinstance(body, list):
        lines = [f"- {item}" for item in body if str(item).strip()]
    elif isinstance(body, str):
        lines = [body]
    else:
        lines = []
    slogan = slots.get("slogan", "")
    if slogan:
        lines.append("")
        lines.append(f"\"{slogan}\"")
    return lines


def _add_fallback_slide(prs: Presentation, slide_data: Dict[str, Any], rendered: Dict[str, Any]) -> None:
    """One slide using simple text boxes (shared styling from rendered)."""
    css = rendered.get("css_variables", {})
    text_color = _hex_to_rgb(css.get("--color-text", "#0F172A"))
    accent_color = _hex_to_rgb(css.get("--color-accent", "#2563EB"))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slots = slide_data.get("slots", {})

    title_box = slide.shapes.add_textbox(Pt(36), Pt(24), Pt(860), Pt(70))
    title_frame = title_box.text_frame
    title_run = title_frame.paragraphs[0].add_run()
    title_run.text = str(slots.get("title", "Untitled Slide"))
    title_run.font.size = Pt(34)
    title_run.font.bold = True
    title_run.font.color.rgb = accent_color

    subtitle_box = slide.shapes.add_textbox(Pt(38), Pt(98), Pt(850), Pt(40))
    subtitle_frame = subtitle_box.text_frame
    subtitle_run = subtitle_frame.paragraphs[0].add_run()
    subtitle_run.text = str(slots.get("subtitle", ""))
    subtitle_run.font.size = Pt(18)
    subtitle_run.font.color.rgb = text_color

    body_box = slide.shapes.add_textbox(Pt(48), Pt(160), Pt(820), Pt(320))
    body_frame = body_box.text_frame
    lines = _as_lines(slots)
    if lines:
        body_frame.paragraphs[0].text = lines[0]
        body_frame.paragraphs[0].font.size = Pt(20)
        body_frame.paragraphs[0].font.color.rgb = text_color
        for line in lines[1:]:
            p = body_frame.add_paragraph()
            p.text = line
            p.font.size = Pt(20)
            p.font.color.rgb = text_color


def _compile_fallback_placeholder_deck(rendered: Dict[str, Any], output_path: str) -> str:
    """Simple geometric layout for all slides."""
    prs = Presentation()
    for slide_data in rendered.get("slides", []):
        _add_fallback_slide(prs, slide_data, rendered)

    prs.save(output_path)
    return output_path


def compile_rendered_to_pptx(rendered: Dict[str, Any], output_path: str) -> str:
    """
    Per slide: if `template_pptx_path` exists, clone that master and inject copy + theme;
    otherwise use simple placeholder boxes. Mixed decks are supported in one export.
    """
    slides = rendered.get("slides", [])
    if not slides:
        raise ValueError("No slides to compile")

    prs = Presentation()
    first_tpl: Optional[str] = None
    for s in slides:
        tp = s.get("template_pptx_path")
        if tp and os.path.isfile(str(tp)):
            first_tpl = str(tp)
            break
    if first_tpl:
        src0 = Presentation(first_tpl)
        prs.slide_width = src0.slide_width
        prs.slide_height = src0.slide_height

    for slide_data in slides:
        tp = slide_data.get("template_pptx_path")
        tidx = int(slide_data.get("template_slide_index", 0))
        if tp and os.path.isfile(str(tp)):
            add_slide_from_template_file(
                prs, str(tp), tidx, slide_data.get("slots", {}), rendered
            )
        else:
            _add_fallback_slide(prs, slide_data, rendered)

    prs.save(output_path)
    return output_path

