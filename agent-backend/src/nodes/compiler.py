from typing import Dict, Any
from ..graph.state import GraphState
from pptx import Presentation
import os

def compiler_node(state: GraphState) -> Dict[str, Any]:
    print("--- COMPILER ---")
    
    template_path = state.get("template_path")
    final_slides = state.get("final_slides_data", [])
    
    print(f"Compiling presentation based on template {template_path}")
    
    for i, slide_data in enumerate(final_slides):
        entity = slide_data['entity']
        data = slide_data['data']
        mapped_fields = data.get("mapped_fields", [])
        
        # Load the presentation template
        prs = Presentation(template_path)
        slide = prs.slides[0] # Working on the assumption template is the first slide
        
        # Helper to find shapes recursively including grouped shapes
        def find_shape_by_id(shapes, shape_id):
            for shape in shapes:
                if shape.shape_id == shape_id:
                    return shape
                if shape.shape_type == 6: # MSO_SHAPE_TYPE.GROUP
                    found = find_shape_by_id(shape.shapes, shape_id)
                    if found:
                        return found
            return None
            
        print(f" - Replacing text for entity: {entity}")
        
        for mapped_field in mapped_fields:
            target_shape_id = mapped_field.get("shape_id")
            new_text = mapped_field.get("new_text", "")
            
            # Find the shape with this ID
            shape = find_shape_by_id(slide.shapes, target_shape_id)
            if shape and shape.has_text_frame:
                 shape.text = new_text
                 print(f"   Injected '{new_text[:20]}...' into shape {target_shape_id}")
                    
        # Apply Graphical Layout Automation (Move checkmarks to scores)
        mapped_visual_elements = data.get("mapped_visual_elements", [])
        if mapped_visual_elements:
            print(" - Applying graphical layout automation...")
            for element in mapped_visual_elements:
                shape_id_to_move = element.get("shape_id_to_move")
                target_shape_id = element.get("target_shape_id")
                
                target_shape = find_shape_by_id(slide.shapes, target_shape_id)
                shape_to_move = find_shape_by_id(slide.shapes, shape_id_to_move)
                        
                if target_shape and shape_to_move:
                    # Calculate center coordinates of the target shape
                    # And align the moving shape to be centered there
                    target_center_x = target_shape.left + (target_shape.width / 2)
                    target_center_y = target_shape.top + (target_shape.height / 2)
                    
                    new_left = target_center_x - (shape_to_move.width / 2)
                    new_top = target_center_y - (shape_to_move.height / 2)
                    
                    # Physically apply coordinate drift to the presentation object
                    shape_to_move.left = int(new_left)
                    shape_to_move.top = int(new_top)
                    
                    print(f"   Moved shape {shape_id_to_move} to align with target shape {target_shape_id}")
                else:
                    print(f"   Warning: Could not find shape {shape_id_to_move} or target {target_shape_id}")
                    
        # Save the new presentation
        output_filename = f"output_{entity.replace(' ', '_')}.pptx"
        # Always output to the root of the project so the user can easily find it
        output_filepath = os.path.join(
            os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__)))), 
            output_filename
        )
        
        prs.save(output_filepath)
        print(f"   => Saved compiled presentation to: {output_filepath}")
        
    print("\nCompiler node finished successfully.")
    
    return {}
