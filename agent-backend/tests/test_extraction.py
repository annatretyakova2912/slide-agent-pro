import os
from pptx import Presentation
import sys
import json

# Add parent directory to path to allow importing from src
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
from src.services.python_pptx import export_slide_context_for_llm

def create_test_presentation():
    prs = Presentation()
    slide_layout = prs.slide_layouts[0] # title slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Hello, World!"
    subtitle.text = "python-pptx was here!"
    
    # save it
    test_path = "test_presentation.pptx"
    prs.save(test_path)
    return test_path

try:
    pptx_path = create_test_presentation()
    print("Created test pptx")
    
    # Now try to extract
    json_result = export_slide_context_for_llm(pptx_path, 0)
    print("Extraction successful. Output length:", len(json_result))
    
    # Verify it parses
    parsed = json.loads(json_result)
    print("Parsed JSON title:", [s['shape_details']['text_content'] for s in parsed['shapes'] if 'text_content' in s['shape_details']])
    
    # clean up
    os.remove(pptx_path)

except Exception as e:
    print("Error:", e)
